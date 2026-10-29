from pathlib import Path

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfgen import canvas

from aigc_selfcheck.docio import read_document, read_docx, read_pdf, read_plaintext


def _make_pdf(path: Path, lines: list[str]) -> None:
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    pdf = canvas.Canvas(str(path))
    pdf.setFont("STSong-Light", 12)
    y = 760
    for line in lines:
        pdf.drawString(72, y, line)
        y -= 24
    pdf.save()


def test_title_and_author_styles_are_skipped(tmp_path: Path) -> None:
    doc = Document()
    doc.styles.add_style("Author", WD_STYLE_TYPE.PARAGRAPH)
    doc.add_paragraph("论文标题", style="Title")
    doc.add_paragraph("张三", style="Author")
    doc.add_paragraph("这是正文段落。", style="Normal")
    path = tmp_path / "sample.docx"
    doc.save(path)

    blocks = read_docx(path)

    assert blocks[0].kind == "heading"
    assert blocks[0].text == "论文标题"
    assert blocks[1].kind == "metadata"
    assert blocks[1].text == "张三"
    assert blocks[2].kind == "paragraph"


def _docx_with_paragraphs(tmp_path: Path, items: list[tuple[str, str]]) -> Path:
    doc = Document()
    existing = {style.name for style in doc.styles}
    for text, style in items:
        if style not in existing:
            doc.styles.add_style(style, WD_STYLE_TYPE.PARAGRAPH)
            existing.add(style)
        doc.add_paragraph(text, style=style)
    path = tmp_path / "handbuilt.docx"
    doc.save(path)
    return path


def test_manual_normal_headings_are_skipped(tmp_path: Path) -> None:
    """真实论文常用 Normal 样式手打标题（目录 toc 样式、标题靠编号/字号）。"""
    path = _docx_with_paragraphs(
        tmp_path,
        [
            ("目  录", "Normal"),
            ("1 总论\t10", "toc 1"),
            ("1.1 项目概况\t10", "toc 2"),
            ("1 总论", "Normal"),
            ("1.1 项目概况", "Normal"),
            ("项目名称为武汉市汉阳区四新“云璟”住宅小区建设项目，拟建于武汉市汉阳区四新片区。", "Normal"),
        ],
    )
    blocks = read_docx(path)

    assert [b.kind for b in blocks] == ["heading", "toc", "toc", "heading", "heading", "paragraph"]


def test_cover_fields_are_metadata(tmp_path: Path) -> None:
    path = _docx_with_paragraphs(
        tmp_path,
        [
            ("题    目\u3000某小区建设项目可行性研究", "Normal"),
            ("学    院\u3000经济与管理学院", "Normal"),
            ("【关键词】\u3000房地产项目\u3000可行性研究", "Normal"),
            ("摘  要", "Normal"),
            ("这是摘要正文段落，长度足够说明它应当被当作正文处理。", "Normal"),
        ],
    )
    blocks = read_docx(path)

    assert [b.kind for b in blocks] == ["metadata", "metadata", "metadata", "heading", "paragraph"]


def test_english_paragraph_is_not_chinese_detectable(tmp_path: Path) -> None:
    """英文摘要不适用中文检测器，应被 classify_block 标为 non_zh。"""
    from aigc_selfcheck.preprocess import classify_block, detectable

    path = _docx_with_paragraphs(
        tmp_path,
        [("As the real estate industry enters a period of deep adjustment, investment decisions become harder.", "Normal")],
    )
    block = classify_block(read_docx(path)[0])

    assert block.kind == "non_zh"
    assert not detectable(block)


def test_body_short_paragraph_without_number_is_detectable(tmp_path: Path) -> None:
    """无编号、无句读的正文短句不能被误判成标题。"""
    path = _docx_with_paragraphs(
        tmp_path,
        [("项目总投资约二十九亿元", "Normal"), ("项目建成后将显著改善区域居住条件并带动就业增长。", "Normal")],
    )
    blocks = read_docx(path)

    assert blocks[0].kind == "paragraph"
    assert blocks[1].kind == "paragraph"


class TestPdfStructure:
    """PDF 没有样式信息，全靠文本启发式——这里覆盖封面/目录/标题。"""

    def test_cover_fields_and_toc(self) -> None:
        from aigc_selfcheck.docio import _pdf_text_to_blocks

        text = (
            "湖北大学知行学院\n本科毕业论文（设计）\n"
            "题    目　某小区建设项目可行性研究\n"
            "学    院　经济与管理学院\n"
            "目  录\n绪论.......7\n1.1 项目概况.......10\n"
            "1 总论\n1.1 项目概况\n"
            "项目名称为某小区建设项目，拟建于武汉市汉阳区。\n"
            "致  谢\n"
        )
        kinds = [b.kind for b in _pdf_text_to_blocks(text)]

        assert kinds.count("toc") == 2
        assert "metadata" in kinds
        assert kinds[-1] == "heading"  # 致  谢
        assert kinds.count("paragraph") >= 1

    def test_leading_cover_demoted_only_in_long_docs(self) -> None:
        from aigc_selfcheck.docio import Block, _demote_leading_short_blocks

        long_doc = [Block(kind="paragraph", text="湖北大学知行学院", index=0)]
        long_doc += [Block(kind="paragraph", text=f"正文第{i}段内容足够长。", index=i + 1) for i in range(9)]
        assert _demote_leading_short_blocks(long_doc)[0].kind == "metadata"

        short_doc = [Block(kind="paragraph", text="短文本第一行", index=0), Block(kind="paragraph", text="正文。", index=1)]
        assert _demote_leading_short_blocks(short_doc)[0].kind == "paragraph"


class TestPlaintextStructure:
    def test_markdown_table_and_list_skipped(self, tmp_path: Path) -> None:
        path = tmp_path / "doc.md"
        path.write_text(
            "# 1 总论\n\n"
            "- [绪论](#绪论)\n- [概况](#概况)\n\n"
            "正文段落内容在这里，描述实验设置与评价指标。\n\n"
            "| 序号 | 费用 | 金额 |\n|---|---|---|\n| 1 | 土地 | 149832 |\n",
            encoding="utf-8",
        )
        blocks = read_plaintext(path)
        kinds = [b.kind for b in blocks]

        assert "list" in kinds
        assert "table" in kinds
        assert kinds.count("heading") == 1

    def test_ordered_short_line_is_heading_not_list(self, tmp_path: Path) -> None:
        path = tmp_path / "doc.md"
        path.write_text(
            "1. 研究背景\n\n正文段落内容在这里，描述研究背景与意义。\n",
            encoding="utf-8",
        )
        blocks = read_plaintext(path)

        assert blocks[0].kind == "heading"

    def test_ordered_long_line_is_list(self, tmp_path: Path) -> None:
        path = tmp_path / "doc.md"
        path.write_text(
            "1. 首先采集样本，覆盖 2015 至 2021 年的论文全文并核对作者信息。\n"
            "2. 然后清洗格式，去除页眉页脚与参考文献。\n",
            encoding="utf-8",
        )
        blocks = read_plaintext(path)

        assert blocks[0].kind == "list"

    def test_plaintext_toc_and_cover(self, tmp_path: Path) -> None:
        path = tmp_path / "doc.txt"
        path.write_text(
            "湖北大学知行学院\n本科毕业论文（设计）\n"
            "题    目　某小区建设项目可行性研究\n"
            "目  录\n绪论.......7\n"
            "项目名称为某小区建设项目，拟建于武汉市汉阳区。\n"
            "致  谢\n",
            encoding="utf-8",
        )
        blocks = read_plaintext(path)
        kinds = [b.kind for b in blocks]

        assert "toc" in kinds
        assert "metadata" in kinds
        assert kinds[-1] == "heading"


def test_pdf_heading_and_paragraph(tmp_path: Path) -> None:
    path = tmp_path / "sample.pdf"
    _make_pdf(
        path,
        [
            "1 引言",
            "这是第一段正文，用来测试 PDF 文本提取。",
            "S = w1 * S_w + w2 * S_k",
        ],
    )

    blocks = read_pdf(path)

    assert blocks[0].kind == "heading"
    assert blocks[0].text == "1 引言"
    assert blocks[1].kind == "paragraph"
    assert "这是第一段正文" in blocks[1].text
    assert blocks[2].kind == "paragraph"
    assert "S = w1" in blocks[2].text


def test_read_document_dispatches_pdf(tmp_path: Path) -> None:
    path = tmp_path / "sample.pdf"
    _make_pdf(path, ["摘 要", "本文提出一种方法。"])

    blocks = read_document(path)

    assert [block.kind for block in blocks] == ["heading", "paragraph"]


def test_pdf_without_text_raises_clear_error(tmp_path: Path) -> None:
    path = tmp_path / "blank.pdf"
    pdf = canvas.Canvas(str(path))
    pdf.save()

    with pytest.raises(ValueError, match="没有可提取文本"):
        read_pdf(path)
